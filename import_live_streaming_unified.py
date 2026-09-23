"""
IMPORT LIVE STREAMING - UNIFIED (Ellips & Cussons)
================================================================================
Satu script buat kedua client. Mapping kolom LIVE STREAMING-nya SAMA PERSIS
antara Ellips & Cussons (B=Creator, D=Room ID, E=Count formula, F=Live Date,
J/K/L=Date/Month/Year formula) -- bedanya cuma nama tab & Cussons punya kolom
"Cek" (C) tambahan yang VLOOKUP ke roster creator, TAPI itu formula otomatis,
script gak perlu nulis situ.

ELLIPS juga diisi kolom Views / Like / Comment / GMV / Product (lihat METRICS)
-- posisi kolomnya dicari dari nama header di tab, bukan huruf yang di-hardcode.
Kalau header di sheet / file sumber gak ketemu, script berhenti sebelum nulis.

CARA PAKAI:
1. Isi TARGET_SPREADSHEET_ID (pilih salah satu dari TARGET_PROFILES) dan
   SOURCE_XLSX_PATH di bawah.
2. python import_live_streaming_unified.py
"""

import datetime
from pathlib import Path

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

# ============== KONFIGURASI -- ISI INI ==============
CREDENTIALS_FILE = 'credentials.json'

TARGET_SPREADSHEET_ID = ''  # Cussons
# TARGET_SPREADSHEET_ID = ''  # Ellips

SOURCE_XLSX_PATH = r'C:\Users\Subhan\OneDrive\Documents\Automasi\TAP_LS Cussons (1-30 August).xlsx'
SOURCE_SHEET_NAME = 'Custom report'

FILTER_YEAR = 2026
FILTER_MONTH = 8  # 8 = Agustus. Set None kalau mau semua bulan.
# ======================================================

SCOPES = ['https://www.googleapis.com/auth/spreadsheets']

TARGET_PROFILES = {
    '1tIG9FhUogXwBJK6YuzpT19nFlJs5EDfXA6EYQ493paE': {
        'name': 'ELLIPS',
        'sheet': 'Live Streaming',
        'has_live_start_end': True,  # G=Live Start, H=End, I=Duration
        'extra_metrics': True,  # Views/Like/Comment/GMV/Product (lihat METRICS)
    },
    '1ZBOvn5fReBECSzgXrAS6SNuq7tWDaem9Cf_QhQN4b_8': {
        'name': 'CUSSONS',
        'sheet': 'LIVE STREAMING',
        'has_live_start_end': True,  # G=Live Start, H=End, I=Duration
        'extra_metrics': False,
    },
}

# Kolom metrik tambahan. Posisi kolom tujuan DICARI dari nama header di baris
# header tab Live Streaming (baris yang sama dengan "Live Room ID"), jadi gak
# tergantung huruf kolom. Nama dicocokkan case-insensitive.
#   key: (kandidat header di sheet tujuan, kandidat header di file TAP_LS, agregasi)
# Agregasi kalau 1 room muncul di beberapa baris file sumber:
#   'max' = metrik level room yang keulang tiap baris (views/like/comment)
#   'sum' = dijumlah per baris unik (Date+Product) -> GMV
#   'list' = gabung nama produk unik, dipisah koma (angka -> ambil max)
METRICS = {
    'views': (['views', 'view', 'live views', 'viewers'],
              ['LIVE views', 'Views', 'Viewers'], 'max'),
    'likes': (['like', 'likes', 'live likes'],
              ['LIVE likes', 'Likes'], 'max'),
    'comments': (['comment', 'comments', 'live comments'],
                 ['LIVE comments', 'Comments'], 'max'),
    'gmv': (['gmv', 'live gmv'],
            ['Creator LIVE-attributed GMV', 'Affiliate LIVE GMV', 'LIVE GMV', 'GMV'], 'sum'),
    'product': (['product', 'products', 'produk', 'product name'],
                ['Product name', 'Products', 'Product'], 'list'),
}

# Mapping kolom -- sama untuk semua target (kalau ada client baru dgn mapping
# beda, tinggal tambah override per-profile di sini).
COLS = {'creator': 'B', 'roomid': 'D', 'count': 'E', 'livedate': 'F',
        'livestart': 'G', 'liveend': 'H', 'duration': 'I',
        'day': 'J', 'month': 'K', 'year': 'L'}


def get_target_profile():
    profile = TARGET_PROFILES.get(TARGET_SPREADSHEET_ID)
    if profile is None:
        raise ValueError(
            f'TARGET_SPREADSHEET_ID "{TARGET_SPREADSHEET_ID}" belum ada di TARGET_PROFILES. '
            'Tambahin dulu profilnya (nama tab Live Streaming) sebelum run.'
        )
    return profile


def get_sheets_service():
    creds = service_account.Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
    return build('sheets', 'v4', credentials=creds)


def extract_year_month(value):
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.year, value.month
    if isinstance(value, str):
        s = value.strip()
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y'):
            try:
                dt = datetime.datetime.strptime(s[:19] if ' ' in s else s, fmt)
                return dt.year, dt.month
            except ValueError:
                continue
    return None


def parse_live_time_info(s):
    """'LIVE time info' formatnya 2 datetime nempel jadi satu string, dipisah 1 strip,
    masing-masing PERSIS 19 karakter ('YYYY-MM-DD HH:MM:SS'), misal:
    '2026-08-19 11:06:52-2026-08-19 14:02:18' -> (start, end)."""
    if not s or not isinstance(s, str):
        return None, None
    s = s.strip()
    if len(s) < 39:
        return None, None
    start_str, end_str = s[:19], s[20:39]
    try:
        start_dt = datetime.datetime.strptime(start_str, '%Y-%m-%d %H:%M:%S')
        end_dt = datetime.datetime.strptime(end_str, '%Y-%m-%d %H:%M:%S')
        return start_dt, end_dt
    except ValueError:
        return None, None


def time_fraction(dt):
    """Konversi jam:menit:detik jadi pecahan hari (format TIME asli Google Sheets)."""
    if dt is None:
        return ''
    return (dt.hour * 3600 + dt.minute * 60 + dt.second) / 86400


def duration_fraction(start_dt, end_dt):
    if start_dt is None or end_dt is None:
        return ''
    delta = end_dt - start_dt
    return delta.total_seconds() / 86400


def to_number(v):
    """Angka dari export TikTok: bisa int/float, 'Rp1.234.567', '1,234', '12.5'."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('Rp', '').replace(' ', '')
    if s in ('', '-', '--'):
        return 0.0
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif s.count('.') > 1 or (s.count('.') == 1 and len(s.split('.')[1]) == 3):
        s = s.replace('.', '')
    elif ',' in s and len(s.split(',')[-1]) == 3:
        s = s.replace(',', '')
    else:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def is_number_like(v):
    if isinstance(v, (int, float)):
        return True
    s = str(v).strip().replace('Rp', '').replace('.', '').replace(',', '')
    return s.isdigit()


def normalize_header(name):
    return ' '.join(str(name or '').split()).lower()


def resolve_source_metric_cols(header):
    """Return {metric_key: index kolom di file sumber}. Error kalau ada yang gak ketemu."""
    by_norm = {}
    for i, h in enumerate(header):
        by_norm.setdefault(normalize_header(h), i)
    cols, missing = {}, []
    for key, (_, src_names, _) in METRICS.items():
        idx = next((by_norm[normalize_header(n)] for n in src_names if normalize_header(n) in by_norm), None)
        if idx is None:
            missing.append(' / '.join(src_names))
        else:
            cols[key] = idx
    if missing:
        raise ValueError(
            f'Kolom metrik berikut tidak ketemu di file sumber: {"; ".join(missing)}.\n'
            f'Header yang ada: {header}\n'
            'Tambahin nama header yang bener ke METRICS di script ini.'
        )
    return cols


def finalize_metrics(acc):
    out = {}
    for key, (_, _, agg) in METRICS.items():
        vals = acc.get(key, [])
        if agg == 'max':
            out[key] = max((to_number(v) for v in vals), default='')
        elif agg == 'sum':
            # dijumlah per baris unik (Date+Product) biar baris dobel persis gak kehitung 2x
            lines = acc.get(f'{key}__lines', {})
            out[key] = sum(to_number(v) for v in lines.values()) if lines else ''
        else:
            vals = [v for v in vals if v not in (None, '', '-')]
            if vals and all(is_number_like(v) for v in vals):
                out[key] = max(to_number(v) for v in vals)
            else:
                uniq = list(dict.fromkeys(str(v).strip() for v in vals))
                out[key] = ', '.join(uniq)
    return out


def read_source_file(path, profile):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if SOURCE_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f'Sheet "{SOURCE_SHEET_NAME}" tidak ketemu. Sheet yang ada: {wb.sheetnames}')
    ws = wb[SOURCE_SHEET_NAME]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError('Sheet sumber kosong.')
    header = [str(h).strip() if h is not None else '' for h in rows[0]]

    def col_idx(name):
        return header.index(name) if name in header else -1

    i_creator = col_idx('Creator name')
    i_roomid = col_idx('Livestream room ID')
    i_livedate = col_idx('LIVE time info')

    missing = [n for n, i in [('Creator name', i_creator), ('Livestream room ID', i_roomid),
                               ('LIVE time info', i_livedate)] if i == -1]
    if missing:
        raise ValueError(f'Kolom berikut tidak ketemu di header: {", ".join(missing)}')

    metric_cols = resolve_source_metric_cols(header) if profile['extra_metrics'] else {}
    i_product = metric_cols.get('product')

    def cell(row, i):
        return row[i] if i is not None and i < len(row) else None

    data_rows = rows[1:]
    by_room = {}
    records = []
    dup_in_source = 0
    filtered_out_month = 0
    skipped_rooms = set()
    for row in data_rows:
        room_id_raw = row[i_roomid]
        date_cell = row[0]
        if not room_id_raw or room_id_raw == '-' or date_cell == 'Summary':
            continue
        room_id = str(room_id_raw).strip()
        if room_id in skipped_rooms:
            continue
        rec = by_room.get(room_id)
        if rec is not None:
            dup_in_source += 1
        else:
            live_date = row[i_livedate]
            if FILTER_MONTH is not None:
                ym = extract_year_month(live_date)
                if ym != (FILTER_YEAR, FILTER_MONTH):
                    filtered_out_month += 1
                    skipped_rooms.add(room_id)
                    continue
            live_start_dt, live_end_dt = parse_live_time_info(live_date)
            rec = {
                'creator': str(row[i_creator] or '').strip(),
                'room_id': room_id,
                'live_date': live_date,
                'live_start_dt': live_start_dt,
                'live_end_dt': live_end_dt,
                '_acc': {},
            }
            by_room[room_id] = rec
            records.append(rec)

        # kumpulin metrik dari SEMUA baris room ini (1 room bisa beberapa baris produk)
        acc = rec['_acc']
        line_key = (str(date_cell), str(cell(row, i_product)))
        for key, i in metric_cols.items():
            val = cell(row, i)
            acc.setdefault(key, []).append(val)
            if METRICS[key][2] == 'sum':
                acc.setdefault(f'{key}__lines', {})[line_key] = val

    for rec in records:
        rec['metrics'] = finalize_metrics(rec.pop('_acc')) if metric_cols else {}

    return records, dup_in_source, filtered_out_month


MONTH_ABBR_ID = {
    1: 'JAN', 2: 'FEB', 3: 'MAR', 4: 'APR', 5: 'MEI', 6: 'JUN',
    7: 'JUL', 8: 'AGUS', 9: 'SEPT', 10: 'OKT', 11: 'NOV', 12: 'DES',
}
MONTH_NAME_EN = {
    1: 'January', 2: 'February', 3: 'March', 4: 'April', 5: 'May', 6: 'June',
    7: 'July', 8: 'August', 9: 'September', 10: 'October', 11: 'November', 12: 'December',
}


def roster_sheet_name(brand, year, month):
    """Nama tab roster creator per bulan -- beda pattern per brand.
    ELLIPS: 'GMV Creator [AGUS]' / CUSSONS: "Creator Performance PZ Cussons August'26"."""
    if brand == 'ELLIPS':
        return f'GMV Creator [{MONTH_ABBR_ID[month]}]'
    return f"Creator Performance PZ Cussons {MONTH_NAME_EN[month]}'{str(year)[-2:]}"


def sync_creator_roster(service, spreadsheet_id, sheet_name, creator_names):
    """TAP-only: creator di TAP itu cuma creator yang kita handle sendiri (beda dari
    SC yang gabungan sama creator luar), jadi kalau ada creator baru yang belum
    tercatat di roster bulan itu berarti roster-nya belum di-update -- tambahin
    otomatis ke kolom B baris kosong paling bawah."""
    col = 'B'
    try:
        result = service.spreadsheets().values().get(
            spreadsheetId=spreadsheet_id, range=f"'{sheet_name}'!{col}:{col}"
        ).execute()
    except HttpError as e:
        raise ValueError(f'Tab roster "{sheet_name}" tidak ketemu / gagal dibaca: {e}')
    values = result.get('values', [])

    existing = set()
    last_filled = 0
    for i, row in enumerate(values):
        v = row[0].strip() if row and row[0] else ''
        if not v or v.lower() == 'creatorusername':
            continue
        existing.add(v.lower())
        last_filled = i + 1

    seen = set()
    missing = []
    for name in creator_names:
        name = (name or '').strip()
        if not name or name.lower() in seen:
            continue
        seen.add(name.lower())
        if name.lower() not in existing:
            missing.append(name)

    if not missing:
        print(f'  Semua creator sudah ada di roster "{sheet_name}".')
        return

    start_row = last_filled + 1
    data = [{'range': f"'{sheet_name}'!{col}{start_row + i}", 'values': [[name]]}
            for i, name in enumerate(missing)]
    service.spreadsheets().values().batchUpdate(
        spreadsheetId=spreadsheet_id, body={'valueInputOption': 'USER_ENTERED', 'data': data}
    ).execute()
    print(f'  {len(missing)} creator baru ditambahin ke roster "{sheet_name}" '
          f'(baris {start_row}-{start_row + len(missing) - 1}): {", ".join(missing)}')


def sync_roster_from_records(service, spreadsheet_id, brand, records_with_date):
    """records_with_date: list of (creator_name, date_value). Dikelompokin per bulan,
    lalu di-sync ke tab roster masing-masing bulan."""
    groups = {}
    for creator, date_val in records_with_date:
        ym = extract_year_month(date_val)
        if ym is None:
            continue
        groups.setdefault(ym, []).append(creator)

    for (year, month), creators in groups.items():
        sheet_name = roster_sheet_name(brand, year, month)
        print(f'Cek roster creator TAP di "{sheet_name}"...')
        sync_creator_roster(service, spreadsheet_id, sheet_name, creators)


def get_existing_room_ids(service, profile):
    sheet = profile['sheet']
    result = service.spreadsheets().values().get(
        spreadsheetId=TARGET_SPREADSHEET_ID,
        range=f"'{sheet}'!{COLS['roomid']}:{COLS['roomid']}"
    ).execute()
    values = result.get('values', [])

    header_row = None
    for i, row in enumerate(values):
        if row and row[0] == 'Live Room ID':
            header_row = i
            break
    if header_row is None:
        raise ValueError(f'Header "Live Room ID" tidak ketemu di kolom {COLS["roomid"]} sheet "{sheet}".')

    data_start = header_row + 1
    existing_ids = set()
    last_filled = data_start - 1
    for i in range(data_start, len(values)):
        rid = values[i][0] if values[i] else None
        if rid:
            existing_ids.add(str(rid).strip())
            last_filled = i

    return existing_ids, data_start + 1, last_filled + 2


def num_to_col_letter(n):
    s = ''
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def get_metric_target_cols(service, profile, header_row_num):
    """Cari huruf kolom Views/Like/Comment/GMV/Product dari baris header tab tujuan."""
    sheet = profile['sheet']
    result = service.spreadsheets().values().get(
        spreadsheetId=TARGET_SPREADSHEET_ID, range=f"'{sheet}'!{header_row_num}:{header_row_num}"
    ).execute()
    header = (result.get('values') or [[]])[0]
    by_norm = {}
    for i, h in enumerate(header):
        by_norm.setdefault(normalize_header(h), i)

    cols, missing = {}, []
    for key, (target_names, _, _) in METRICS.items():
        idx = next((by_norm[n] for n in target_names if n in by_norm), None)
        if idx is None:
            missing.append(' / '.join(target_names))
        else:
            cols[key] = num_to_col_letter(idx + 1)
    if missing:
        raise ValueError(
            f'Header kolom berikut tidak ketemu di baris {header_row_num} tab "{sheet}": {"; ".join(missing)}.\n'
            f'Header yang ada: {header}'
        )
    return cols


def to_serial_date(dt):
    if dt is None:
        return ''
    if isinstance(dt, str):
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try:
                dt = datetime.datetime.strptime(dt[:19], fmt)
                break
            except ValueError:
                continue
        else:
            return dt
    base = datetime.date(1899, 12, 30)
    d = dt.date() if isinstance(dt, datetime.datetime) else dt
    return (d - base).days


def get_sheet_id(service, sheet_name):
    meta = service.spreadsheets().get(spreadsheetId=TARGET_SPREADSHEET_ID, fields='sheets.properties').execute()
    for sh in meta['sheets']:
        if sh['properties']['title'] == sheet_name:
            return sh['properties']['sheetId']
    raise ValueError(f'Sheet "{sheet_name}" tidak ketemu.')


def col_letter_to_num(letters):
    n = 0
    for c in letters:
        n = n * 26 + (ord(c) - 64)
    return n


def ensure_date_time_formats(service, profile, next_row, last_row):
    """Pastiin kolom Live Date ke-format DATE, dan Live Start/End/Duration (kalau ada)
    ke-format TIME -- biar gak ke-render angka mentah kayak '46182'."""
    sheet_id = get_sheet_id(service, profile['sheet'])
    requests = []

    def fmt_request(col_letter, fmt_type, pattern):
        c = col_letter_to_num(col_letter) - 1
        requests.append({
            'repeatCell': {
                'range': {'sheetId': sheet_id, 'startRowIndex': next_row - 1, 'endRowIndex': last_row,
                          'startColumnIndex': c, 'endColumnIndex': c + 1},
                'cell': {'userEnteredFormat': {'numberFormat': {'type': fmt_type, 'pattern': pattern}}},
                'fields': 'userEnteredFormat.numberFormat'
            }
        })

    fmt_request(COLS['livedate'], 'DATE', 'yyyy-mm-dd')
    if profile['has_live_start_end']:
        fmt_request(COLS['livestart'], 'TIME', 'hh:mm:ss')
        fmt_request(COLS['liveend'], 'TIME', 'hh:mm:ss')
        fmt_request(COLS['duration'], 'TIME', '[h]:mm:ss')

    service.spreadsheets().batchUpdate(spreadsheetId=TARGET_SPREADSHEET_ID, body={'requests': requests}).execute()


def write_new_records(service, profile, to_insert, next_row, metric_target_cols):
    sheet_ref = f"'{profile['sheet']}'!"
    value_ranges = []

    for idx, rec in enumerate(to_insert):
        r = next_row + idx
        for key, col in metric_target_cols.items():
            value_ranges.append({'range': f"{sheet_ref}{col}{r}", 'values': [[rec['metrics'].get(key, '')]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['creator']}{r}", 'values': [[rec['creator']]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['roomid']}{r}", 'values': [[rec['room_id']]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['livedate']}{r}", 'values': [[to_serial_date(rec['live_date'])]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['count']}{r}", 'values': [[f"=IF({COLS['roomid']}{r}=\"\";0;1)"]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['day']}{r}", 'values': [[f"=IF(${COLS['livedate']}{r}=\"\";\"\";DAY(${COLS['livedate']}{r}))"]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['month']}{r}", 'values': [[f"=IF(${COLS['livedate']}{r}=\"\";\"\";MONTH(${COLS['livedate']}{r}))"]]})
        value_ranges.append({'range': f"{sheet_ref}{COLS['year']}{r}", 'values': [[f"=IF(${COLS['livedate']}{r}=\"\";\"\";YEAR(${COLS['livedate']}{r}))"]]})

        if profile['has_live_start_end']:
            value_ranges.append({'range': f"{sheet_ref}{COLS['livestart']}{r}", 'values': [[time_fraction(rec.get('live_start_dt'))]]})
            value_ranges.append({'range': f"{sheet_ref}{COLS['liveend']}{r}", 'values': [[time_fraction(rec.get('live_end_dt'))]]})
            value_ranges.append({'range': f"{sheet_ref}{COLS['duration']}{r}", 'values': [[duration_fraction(rec.get('live_start_dt'), rec.get('live_end_dt'))]]})

    entries_per_row = len(value_ranges) // len(to_insert)
    chunk_rows = 700
    for start in range(0, len(to_insert), chunk_rows):
        n = min(chunk_rows, len(to_insert) - start)
        chunk = value_ranges[start * entries_per_row: (start + n) * entries_per_row]
        service.spreadsheets().values().batchUpdate(
            spreadsheetId=TARGET_SPREADSHEET_ID,
            body={'valueInputOption': 'USER_ENTERED', 'data': chunk}
        ).execute()
        print(f'  ... tertulis {start + n}/{len(to_insert)} baris')

    ensure_date_time_formats(service, profile, next_row, next_row + len(to_insert) - 1)


def main():
    if not Path(SOURCE_XLSX_PATH).exists():
        print(f'ERROR: file tidak ketemu di path: {SOURCE_XLSX_PATH}')
        return

    profile = get_target_profile()
    print(f'Target terdeteksi: {profile["name"]} (tab "{profile["sheet"]}")')

    print('Membaca file sumber...')
    records, dup_in_source, filtered_out_month = read_source_file(SOURCE_XLSX_PATH, profile)
    print(f'Ketemu {len(records)} livestream unik ({dup_in_source} baris tambahan per room digabung).')
    if FILTER_MONTH is not None:
        print(f'Filter bulan aktif ({FILTER_MONTH}/{FILTER_YEAR}): {filtered_out_month} livestream di luar bulan itu dilewati.')

    print('Connect ke Google Sheets, ambil Live Room ID yang udah ada...')
    service = get_sheets_service()
    existing_ids, data_start_sheet_row, next_row = get_existing_room_ids(service, profile)
    print(f'Ada {len(existing_ids)} Live Room ID yang udah ada di sheet.')

    metric_target_cols = {}
    if profile['extra_metrics']:
        metric_target_cols = get_metric_target_cols(service, profile, data_start_sheet_row - 1)
        print('Kolom metrik: ' + ', '.join(f'{k}={c}' for k, c in metric_target_cols.items()))

    to_insert = [r for r in records if r['room_id'] not in existing_ids]
    skipped = len(records) - len(to_insert)

    if not to_insert:
        print(f'\nTidak ada livestream baru. Semua ({len(records)}) sudah pernah diinput sebelumnya.')
        return

    print(f'\n{len(to_insert)} livestream baru akan ditambahkan ke baris {next_row}-{next_row + len(to_insert) - 1}.')
    print(f'{skipped} dilewati karena sudah ada di sheet.')
    print('Menulis ke Google Sheets...')
    write_new_records(service, profile, to_insert, next_row, metric_target_cols)

    print('\nSinkronisasi roster creator (khusus sumber TAP)...')
    sync_roster_from_records(
        service, TARGET_SPREADSHEET_ID, profile['name'],
        [(r['creator'], r['live_date']) for r in to_insert]
    )

    print(f'\nSELESAI. {len(to_insert)} livestream baru ditambahkan, {skipped} dilewati (duplikat).')


if __name__ == '__main__':
    main()
