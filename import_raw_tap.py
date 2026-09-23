"""
IMPORT TAP_GMV -> RAW TAP (jalan lokal dari komputer, pakai Google Sheets API)
================================================================================
Baca file TAP_GMV langsung dari komputer, isi semua 43 kolom RAW TAP sesuai
nama kolomnya (termasuk kolom formula: Week, cek usn ext, cek usn sheet gmv
june, cek product id, GMV SHARE LINK).

DEDUP RINGAN: kombinasi Date + Creator name + Product ID yang PERSIS sama
dengan yang udah ada di sheet akan dilewati.

TIDAK auto-nambah baris kalau kurang -- kalau baris di sheet gak cukup,
script berhenti & kasih tau berapa baris yang perlu ditambahin manual dulu
(tambahin manual di Google Sheets: klik kanan baris terakhir > Insert X baris).

CARA PAKAI:
1. Isi CREDENTIALS_FILE, TARGET_SPREADSHEET_ID, SOURCE_XLSX_PATH di bawah.
2. Cek nama tab data di file TAP_GMV kamu -- kalau bukan "Custom report"
   (misal "Sheet1"), ganti SOURCE_SHEET_NAME di bawah.
3. python import_raw_tap.py
"""

import datetime
from pathlib import Path

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build

# ============== KONFIGURASI -- ISI INI ==============
CREDENTIALS_FILE = 'credentials.json'
TARGET_SPREADSHEET_ID = ''
SOURCE_XLSX_PATH = r''
SOURCE_SHEET_NAME = 'Custom report'  # cek dulu, kadang namanya "Sheet1"
TARGET_SHEET_NAME = 'RAW TAP'
# ======================================================

SCOPES = ['https://www.googleapis.com/auth/spreadsheets']

# Pemetaan kolom tujuan (RAW TAP) <- nama kolom di file sumber (TAP_GMV).
# Beberapa kolom dikasih beberapa alias karena nama kolomnya suka beda-beda
# antar bulan (misal "Affiliate GMV" vs "Creator-attributed GMV").
COLUMN_MAP = [
    ('A', ['Date']),
    ('C', ['Comparison date']),
    ('D', ['Campaign ID']),
    ('E', ['Campaign name']),
    ('F', ['Campaign duration']),
    ('G', ['Creator name']),
    ('J', ['Product ID']),
    ('L', ['Product name']),
    ('M', ['Shop ID']),
    ('N', ['Shop name']),
    ('O', ['Level 1 category']),
    ('P', ['Level 2 category']),
    ('Q', ['Creator-attributed GMV', 'Affiliate GMV']),
    ('R', ['Affiliate video-attributed GMV', 'Affiliate video GMV']),
    ('S', ['Creator LIVE-attributed GMV', 'Affiliate LIVE GMV']),
    ('U', ['Creator-attributed orders', 'Orders']),
    ('V', ['Creator LIVE-attributed orders', 'LIVE orders']),
    ('W', ['Creator video-attributed orders', 'Video orders']),
    ('X', ['LIVE likes']),
    ('Y', ['Video likes']),
    ('Z', ['Video views']),
    ('AA', ['LIVE views']),
    ('AB', ['LIVE streams']),
    ('AC', ['Videos']),
    ('AD', ['Products added to Showcase']),
    ('AE', ['Estimated affiliate partner commission']),
    ('AF', ['Actual affiliate partner commission']),
    ('AG', ['Estimated creator commission']),
    ('AH', ['Actual creator commission']),
    ('AI', ['GMV (refund)']),
    ('AJ', ['Settled GMV']),
    ('AK', ['Revenue (Showcase)']),
    ('AL', ['Creator-attributed items sold', 'Items sold']),
    ('AM', ['Link GMV']),
    ('AN', ['Link items sold']),
    ('AO', ['Link orders']),
    ('AP', ['Link partner est commission', 'Link partner est. commission']),
    ('AQ', ['Link creator est commission', 'Link creator est. commission']),
]


def get_sheets_service():
    creds = service_account.Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
    return build('sheets', 'v4', credentials=creds)


def find_col(header, names):
    for name in names:
        for i, h in enumerate(header):
            if h == name:
                return i
    return -1


def parse_rupiah(val):
    """Parse angka format 'Rp1.234.567' / 0 / None jadi float."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace('Rp', '').replace('.', '').strip()
    if s in ('', '-', '--'):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def read_source_file(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if SOURCE_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f'Sheet "{SOURCE_SHEET_NAME}" tidak ketemu. Sheet yang ada: {wb.sheetnames}')
    ws = wb[SOURCE_SHEET_NAME]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError('Sheet sumber kosong.')
    header = [str(h).strip() if h is not None else '' for h in rows[0]]

    src_col_index = {}
    missing = []
    for dest_col, names in COLUMN_MAP:
        idx = find_col(header, names)
        if idx == -1:
            missing.append(' / '.join(names))
        else:
            src_col_index[dest_col] = idx
    if missing:
        raise ValueError(f'Kolom berikut tidak ketemu di header: {", ".join(missing)}')

    idx_date = find_col(header, ['Date'])
    idx_creator = find_col(header, ['Creator name'])
    idx_product = find_col(header, ['Product ID'])

    data_rows = rows[1:]
    if len(data_rows) < 4:
        print(f'PERINGATAN: sheet sumber cuma ada {len(data_rows)} baris data. '
              f'Kalau harusnya data 1 bulan penuh, kemungkinan datanya kurang lengkap.')

    valid_rows = [r for r in data_rows if r[idx_date] and r[idx_date] != 'Summary']

    idx_gmv = src_col_index['Q']
    before_gmv_filter = len(valid_rows)
    valid_rows = [r for r in valid_rows if parse_rupiah(r[idx_gmv]) != 0]
    print(f'Filter Affiliate GMV != 0: {before_gmv_filter - len(valid_rows)} baris GMV-nya 0 dibuang, '
          f'sisa {len(valid_rows)} baris.')

    return valid_rows, src_col_index, idx_date, idx_creator, idx_product


def date_to_text(dt):
    """Format tanggal jadi teks yyyy-mm-dd (dibutuhin formula Week yang pakai DATEVALUE)."""
    if dt is None:
        return ''
    if isinstance(dt, datetime.datetime):
        return dt.strftime('%Y-%m-%d')
    if isinstance(dt, datetime.date):
        return dt.strftime('%Y-%m-%d')
    return str(dt)


def get_sheet_row_count(service, spreadsheet_id, sheet_name):
    meta = service.spreadsheets().get(
        spreadsheetId=spreadsheet_id, fields='sheets(properties(title,gridProperties))'
    ).execute()
    for sh in meta['sheets']:
        if sh['properties']['title'] == sheet_name:
            return sh['properties']['gridProperties']['rowCount']
    raise ValueError(f'Sheet "{sheet_name}" tidak ketemu.')


def get_existing_keys(service):
    """Baca kolom A (Date), G (Creator), J (Product ID) di RAW TAP, chunk per 30rb baris."""
    target_max_row = get_sheet_row_count(service, TARGET_SPREADSHEET_ID, TARGET_SHEET_NAME)

    chunk_size = 30000
    col_a, col_g, col_j = [], [], []
    for start in range(1, target_max_row + 1, chunk_size):
        end = min(start + chunk_size - 1, target_max_row)
        resp = service.spreadsheets().values().batchGet(
            spreadsheetId=TARGET_SPREADSHEET_ID,
            ranges=[f"'{TARGET_SHEET_NAME}'!A{start}:A{end}",
                    f"'{TARGET_SHEET_NAME}'!G{start}:G{end}",
                    f"'{TARGET_SHEET_NAME}'!J{start}:J{end}"]
        ).execute()
        vr = resp.get('valueRanges', [])
        a_vals = vr[0].get('values', []) if len(vr) > 0 else []
        g_vals = vr[1].get('values', []) if len(vr) > 1 else []
        j_vals = vr[2].get('values', []) if len(vr) > 2 else []
        col_a.extend(a_vals)
        col_g.extend(g_vals)
        col_j.extend(j_vals)
        if not a_vals and not g_vals and not j_vals:
            break

    header_row = None
    for i, row in enumerate(col_a):
        if row and row[0] == 'Date':
            header_row = i
            break
    if header_row is None:
        raise ValueError('Header "Date" tidak ketemu di kolom A sheet RAW TAP.')

    data_start = header_row + 1
    existing_keys = set()
    last_filled = data_start - 1
    last_row = max(len(col_a), len(col_g), len(col_j))
    for i in range(data_start, last_row):
        d = col_a[i][0] if i < len(col_a) and col_a[i] else ''
        c = col_g[i][0] if i < len(col_g) and col_g[i] else ''
        p = col_j[i][0] if i < len(col_j) and col_j[i] else ''
        if d or c or p:
            last_filled = i
        if d and c:
            existing_keys.add(f'{str(d).strip()}|{str(c).strip()}|{str(p).strip()}')

    return existing_keys, data_start + 1, last_filled + 2, target_max_row


def write_new_rows(service, to_insert, src_col_index, next_row, data_start_sheet_row):
    sheet_ref = f"'{TARGET_SHEET_NAME}'!"
    value_ranges = []

    for idx, row in enumerate(to_insert):
        r = next_row + idx
        for dest_col, _ in COLUMN_MAP:
            val = row[src_col_index[dest_col]]
            if dest_col == 'A':
                val = date_to_text(val)
            value_ranges.append({'range': f'{sheet_ref}{dest_col}{r}', 'values': [[val if val is not None else '']]})

        value_ranges.append({
            'range': f'{sheet_ref}B{r}',
            'values': [[f'=UPPER(TEXT(DATEVALUE(A{r});"mmm"))&" W"&ROUNDUP(DAY(DATEVALUE(A{r}))/7;0)&" ("&'
                        f'((ROUNDUP(DAY(DATEVALUE(A{r}))/7;0)-1)*7+1)&"-"&MIN(ROUNDUP(DAY(DATEVALUE(A{r}))/7;0)*7;'
                        f'DAY(EOMONTH(DATEVALUE(A{r});0)))&")"']]
        })
        value_ranges.append({'range': f'{sheet_ref}H{r}', 'values': [[f"=VLOOKUP(G{r};'GMV Creator JULY'!$B:$B;1;0)"]]})
        value_ranges.append({'range': f'{sheet_ref}I{r}', 'values': [[f"=VLOOKUP(G{r};'GMV Creator JUNE'!$B:$B;1;0)"]]})
        value_ranges.append({'range': f'{sheet_ref}K{r}', 'values': [[f"=VLOOKUP(J{r};'RAW TAP SC'!$E:$E;1;0)"]]})
        value_ranges.append({'range': f'{sheet_ref}T{r}', 'values': [[f'=AK{r}+AM{r}']]})

    entries_per_row = len(COLUMN_MAP) + 5
    chunk_rows = 300
    for start in range(0, len(to_insert), chunk_rows):
        n = min(chunk_rows, len(to_insert) - start)
        chunk = value_ranges[start * entries_per_row: (start + n) * entries_per_row]
        service.spreadsheets().values().batchUpdate(
            spreadsheetId=TARGET_SPREADSHEET_ID,
            body={'valueInputOption': 'USER_ENTERED', 'data': chunk}
        ).execute()
        print(f'  ... tertulis {start + n}/{len(to_insert)} baris')


def main():
    if not Path(SOURCE_XLSX_PATH).exists():
        print(f'ERROR: file tidak ketemu di path: {SOURCE_XLSX_PATH}')
        return

    print('Membaca file sumber...')
    valid_rows, src_col_index, idx_date, idx_creator, idx_product = read_source_file(SOURCE_XLSX_PATH)
    print(f'Ketemu {len(valid_rows)} baris data valid.')
    if not valid_rows:
        print('Tidak ada baris data valid, berhenti.')
        return

    print('Connect ke Google Sheets, ambil data existing buat dedup...')
    service = get_sheets_service()
    existing_keys, data_start_sheet_row, next_row, target_max_row = get_existing_keys(service)
    print(f'Ada {len(existing_keys)} kombinasi Date+Creator+Product ID yang udah ada di sheet.')

    to_insert = []
    skipped = 0
    seen_now = set(existing_keys)
    for row in valid_rows:
        key = f'{str(row[idx_date]).strip()}|{str(row[idx_creator]).strip()}|{str(row[idx_product]).strip()}'
        if key in seen_now:
            skipped += 1
            continue
        seen_now.add(key)
        to_insert.append(row)

    if not to_insert:
        print(f'\nSemua {len(valid_rows)} baris sudah ada sebelumnya (kombinasi Date+Creator+Product ID sama).')
        return

    needed_last_row = next_row + len(to_insert) - 1
    if needed_last_row > target_max_row:
        rows_short = needed_last_row - target_max_row
        print(f'\nERROR: Baris di sheet RAW TAP gak cukup. Butuh sampai baris {needed_last_row}, '
              f'tapi grid sheet cuma sampai baris {target_max_row} (kurang {rows_short} baris).')
        print(f'Tambahin manual dulu minimal {rows_short} baris di bagian bawah sheet RAW TAP '
              f'(klik kanan nomor baris terakhir > Insert {rows_short} rows below), baru jalanin lagi script ini.')
        return

    print(f'\n{len(to_insert)} baris baru akan ditambahkan ke baris {next_row}-{next_row + len(to_insert) - 1}.')
    print(f'{skipped} dilewati (kombinasi Date+Creator+Product ID sudah ada).')
    print('Menulis ke Google Sheets...')
    write_new_rows(service, to_insert, src_col_index, next_row, data_start_sheet_row)

    print(f'\nSELESAI. {len(to_insert)} baris baru ditambahkan, {skipped} dilewati.')


if __name__ == '__main__':
    main()
