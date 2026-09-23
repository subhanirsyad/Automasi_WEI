"""
IMPORT TAP_GMV -> TAP Cussons (1-16 August)  [khusus Cussons]
================================================================================
Beda dari import_raw_tap.py punya Ellips (isi RAW TAP) -- destinasi Cussons ini
kolomnya lebih sedikit & formula-formulanya BEDA (GMV SL, Avg item Sold, Item
Sold LS/SV/SL). Sumbernya (TAP_GMV Cussons) formatnya identik sama TAP_GMV Ellips.

DEDUP: kombinasi Date + Creator name + Product ID yang PERSIS sama dengan yang
udah ada di sheet akan dilewati.

TIDAK auto-nambah baris kalau kurang -- eh sebenernya INI beda dari yang Ellips:
sheet Cussons ini otomatis DITAMBAHIN baris kalau kurang (grid-nya kecil, cuma
1000 baris kapasitas awal).

CARA PAKAI:
1. Isi TARGET_SPREADSHEET_ID, SOURCE_XLSX_PATH di bawah kalau beda.
2. python import_raw_tap_cussons.py
"""

import datetime
from pathlib import Path

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build

# ============== KONFIGURASI -- ISI INI ==============
CREDENTIALS_FILE = 'credentials.json'
TARGET_SPREADSHEET_ID = '1ZBOvn5fReBECSzgXrAS6SNuq7tWDaem9Cf_QhQN4b_8'
SOURCE_XLSX_PATH = r'C:\Users\Subhan\OneDrive\Documents\Automasi\TAP_GMV Cussons (24-30 August).xlsx'
SOURCE_SHEET_NAME = 'Custom report'
TARGET_SHEET_NAME = 'TAP Cussons (1-16  August)'  # PERHATIAN: ada 2 spasi di "1-16  August"
# ======================================================

SCOPES = ['https://www.googleapis.com/auth/spreadsheets']

# Pemetaan kolom tujuan (TAP Cussons) <- nama kolom di file sumber (TAP_GMV).
# Kolom yang GAK ada di sini (WEEK, Cek, GMV SL, Avg. item Sold, Item Sold
# LS/SV/SL) itu FORMULA -- ditulis terpisah di write_formula_cols_().
COLUMN_MAP = [
    ('A', ['Date']),
    ('C', ['Campaign ID']),
    ('D', ['Campaign name']),
    ('E', ['Campaign duration']),
    ('F', ['Creator name']),
    ('H', ['Product ID']),
    ('I', ['Product name']),
    ('J', ['Creator-attributed GMV', 'Affiliate GMV']),
    ('K', ['Affiliate video-attributed GMV', 'Affiliate video GMV']),
    ('L', ['Creator LIVE-attributed GMV', 'Affiliate LIVE GMV']),
    ('N', ['Creator-attributed orders', 'Orders']),
    ('O', ['Creator LIVE-attributed orders', 'LIVE orders']),
    ('P', ['Creator video-attributed orders', 'Video orders']),
    ('Q', ['LIVE likes']),
    ('R', ['Video likes']),
    ('S', ['Video views']),
    ('T', ['LIVE views']),
    ('U', ['LIVE streams']),
    ('V', ['Videos']),
    ('W', ['GMV (refund)']),
    ('X', ['Settled GMV']),
    ('Y', ['Revenue (Showcase)']),
    ('Z', ['Creator-attributed items sold', 'Items sold']),
    ('AA', ['Link GMV']),
    ('AB', ['Link items sold']),
    ('AC', ['Link orders']),
]


def get_sheets_service():
    creds = service_account.Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
    return build('sheets', 'v4', credentials=creds)


def find_col(header, names):
    for name in names:
        for i, h in enumerate(header):
            if h == name:
                return i
    return -1


def parse_rupiah(val):
    """Parse angka format 'Rp1.234.567' / 0 / None jadi float."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace('Rp', '').replace('.', '').strip()
    if s in ('', '-', '--'):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def read_source_file(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if SOURCE_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f'Sheet "{SOURCE_SHEET_NAME}" tidak ketemu. Sheet yang ada: {wb.sheetnames}')
    ws = wb[SOURCE_SHEET_NAME]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError('Sheet sumber kosong.')
    header = [str(h).strip() if h is not None else '' for h in rows[0]]

    src_col_index = {}
    missing = []
    for dest_col, names in COLUMN_MAP:
        idx = find_col(header, names)
        if idx == -1:
            missing.append(' / '.join(names))
        else:
            src_col_index[dest_col] = idx
    if missing:
        raise ValueError(f'Kolom berikut tidak ketemu di header: {", ".join(missing)}')

    idx_date = find_col(header, ['Date'])
    idx_creator = find_col(header, ['Creator name'])
    idx_product = find_col(header, ['Product ID'])

    data_rows = rows[1:]
    valid_rows = [r for r in data_rows if r[idx_date] and r[idx_date] != 'Summary']

    idx_gmv = src_col_index['J']  # 'Creator-attributed GMV'
    before_gmv_filter = len(valid_rows)
    valid_rows = [r for r in valid_rows if parse_rupiah(r[idx_gmv]) != 0]
    print(f'Filter Creator-attributed GMV != 0: {before_gmv_filter - len(valid_rows)} baris GMV-nya 0 dibuang, '
          f'sisa {len(valid_rows)} baris.')

    return valid_rows, src_col_index, idx_date, idx_creator, idx_product


def date_to_text(dt):
    if dt is None:
        return ''
    if isinstance(dt, (datetime.datetime, datetime.date)):
        return dt.strftime('%Y-%m-%d')
    return str(dt)


def get_sheet_id_and_row_count(service):
    meta = service.spreadsheets().get(
        spreadsheetId=TARGET_SPREADSHEET_ID, fields='sheets(properties(sheetId,title,gridProperties))'
    ).execute()
    for sh in meta['sheets']:
        if sh['properties']['title'] == TARGET_SHEET_NAME:
            return sh['properties']['sheetId'], sh['properties']['gridProperties']['rowCount']
    raise ValueError(f'Sheet "{TARGET_SHEET_NAME}" tidak ketemu.')


def ensure_rows(service, rows_needed):
    sheet_id, current_rows = get_sheet_id_and_row_count(service)
    if rows_needed <= 0:
        return current_rows
    service.spreadsheets().batchUpdate(
        spreadsheetId=TARGET_SPREADSHEET_ID,
        body={'requests': [{
            'appendDimension': {'sheetId': sheet_id, 'dimension': 'ROWS', 'length': rows_needed}
        }]}
    ).execute()
    return current_rows + rows_needed


def get_existing_keys(service):
    """Baca kolom A (Date), F (Creator), H (Product ID)."""
    _, target_max_row = get_sheet_id_and_row_count(service)

    resp = service.spreadsheets().values().batchGet(
        spreadsheetId=TARGET_SPREADSHEET_ID,
        ranges=[f"'{TARGET_SHEET_NAME}'!A1:A{target_max_row}",
                f"'{TARGET_SHEET_NAME}'!F1:F{target_max_row}",
                f"'{TARGET_SHEET_NAME}'!H1:H{target_max_row}"]
    ).execute()
    vr = resp.get('valueRanges', [])
    col_a = vr[0].get('values', []) if len(vr) > 0 else []
    col_f = vr[1].get('values', []) if len(vr) > 1 else []
    col_h = vr[2].get('values', []) if len(vr) > 2 else []

    header_row = None
    for i, row in enumerate(col_a):
        if row and row[0] == 'Date':
            header_row = i
            break
    if header_row is None:
        raise ValueError('Header "Date" tidak ketemu di kolom A.')

    data_start = header_row + 1
    existing_keys = set()
    last_filled = data_start - 1
    last_row = max(len(col_a), len(col_f), len(col_h))
    for i in range(data_start, last_row):
        d = col_a[i][0] if i < len(col_a) and col_a[i] else ''
        c = col_f[i][0] if i < len(col_f) and col_f[i] else ''
        p = col_h[i][0] if i < len(col_h) and col_h[i] else ''
        if d or c or p:
            last_filled = i
        if d and c:
            existing_keys.add(f'{str(d).strip()}|{str(c).strip()}|{str(p).strip()}')

    return existing_keys, data_start + 1, last_filled + 2, target_max_row


def write_new_rows(service, to_insert, src_col_index, next_row):
    sheet_ref = f"'{TARGET_SHEET_NAME}'!"
    value_ranges = []

    for idx, row in enumerate(to_insert):
        r = next_row + idx
        for dest_col, _ in COLUMN_MAP:
            val = row[src_col_index[dest_col]]
            if dest_col == 'A':
                val = date_to_text(val)
            value_ranges.append({'range': f'{sheet_ref}{dest_col}{r}', 'values': [[val if val is not None else '']]})

        # kolom formula -- pola persis sama kayak yg udah ada di baris 2-3 sheet ini
        value_ranges.append({
            'range': f'{sheet_ref}B{r}',
            'values': [[f'=UPPER(TEXT(A{r};"MMM"))&" W"&IF(DAY(A{r})<=2;1;IF(DAY(A{r})<=9;2;IF(DAY(A{r})<=16;3;'
                        f'IF(DAY(A{r})<=23;4;5))))&" ("&IF(DAY(A{r})<=2;1;IF(DAY(A{r})<=9;3;IF(DAY(A{r})<=16;10;'
                        f'IF(DAY(A{r})<=23;17;24))))&"-"&IF(DAY(A{r})<=2;2;IF(DAY(A{r})<=9;9;IF(DAY(A{r})<=16;16;'
                        f'IF(DAY(A{r})<=23;23;DAY(EOMONTH(A{r};0))))))&")"']]
        })
        value_ranges.append({
            'range': f'{sheet_ref}G{r}',
            'values': [[f"=VLOOKUP(F{r};'Creator Performance PZ Cussons August''26'!$B:$B;1;0)"]]
        })
        value_ranges.append({'range': f'{sheet_ref}M{r}', 'values': [[f'=AA{r}+Y{r}']]})
        value_ranges.append({'range': f'{sheet_ref}AD{r}', 'values': [[f'=Z{r}/N{r}']]})
        value_ranges.append({'range': f'{sheet_ref}AE{r}', 'values': [[f'=AD{r}*O{r}']]})
        value_ranges.append({'range': f'{sheet_ref}AF{r}', 'values': [[f'=AD{r}*P{r}']]})
        value_ranges.append({'range': f'{sheet_ref}AG{r}', 'values': [[f'=Z{r}-SUM(AE{r};AF{r})']]})

    entries_per_row = len(COLUMN_MAP) + 7
    chunk_rows = 300
    for start in range(0, len(to_insert), chunk_rows):
        n = min(chunk_rows, len(to_insert) - start)
        chunk = value_ranges[start * entries_per_row: (start + n) * entries_per_row]
        service.spreadsheets().values().batchUpdate(
            spreadsheetId=TARGET_SPREADSHEET_ID,
            body={'valueInputOption': 'USER_ENTERED', 'data': chunk}
        ).execute()
        print(f'  ... tertulis {start + n}/{len(to_insert)} baris')


def main():
    if not Path(SOURCE_XLSX_PATH).exists():
        print(f'ERROR: file tidak ketemu di path: {SOURCE_XLSX_PATH}')
        return

    print('Membaca file sumber...')
    valid_rows, src_col_index, idx_date, idx_creator, idx_product = read_source_file(SOURCE_XLSX_PATH)
    print(f'Ketemu {len(valid_rows)} baris data valid.')
    if not valid_rows:
        print('Tidak ada baris data valid, berhenti.')
        return

    print('Connect ke Google Sheets, ambil data existing buat dedup...')
    service = get_sheets_service()
    existing_keys, data_start_sheet_row, next_row, target_max_row = get_existing_keys(service)
    print(f'Ada {len(existing_keys)} kombinasi Date+Creator+Product ID yang udah ada di sheet.')

    to_insert = []
    skipped = 0
    seen_now = set(existing_keys)
    for row in valid_rows:
        key = f'{str(row[idx_date]).strip()}|{str(row[idx_creator]).strip()}|{str(row[idx_product]).strip()}'
        if key in seen_now:
            skipped += 1
            continue
        seen_now.add(key)
        to_insert.append(row)

    if not to_insert:
        print(f'\nSemua {len(valid_rows)} baris sudah ada sebelumnya (kombinasi Date+Creator+Product ID sama).')
        return

    needed_last_row = next_row + len(to_insert) - 1
    if needed_last_row > target_max_row:
        rows_short = needed_last_row - target_max_row
        print(f'\nGrid sheet kurang {rows_short} baris (butuh sampai baris {needed_last_row}, '
              f'tersedia sampai {target_max_row}). Nambah otomatis...')
        target_max_row = ensure_rows(service, rows_short)
        print(f'Baris ditambah, grid sekarang sampai baris {target_max_row}.')

    print(f'\n{len(to_insert)} baris baru akan ditambahkan ke baris {next_row}-{next_row + len(to_insert) - 1}.')
    print(f'{skipped} dilewati (kombinasi Date+Creator+Product ID sudah ada).')
    print('Menulis ke Google Sheets...')
    write_new_rows(service, to_insert, src_col_index, next_row)

    print(f'\nSELESAI. {len(to_insert)} baris baru ditambahkan, {skipped} dilewati.')


if __name__ == '__main__':
    main()
