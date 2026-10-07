"""
INTI - baca daftar creator, cari kontaknya di Kalodata, tulis hasilnya
================================================================================
Input  : Excel ATAU Google Spreadsheet berisi username TikTok creator (atau creator ID).
Output : - Input Excel       -> file Excel hasil (disimpan berkala tiap 25 creator).
         - Input Spreadsheet -> TIDAK ada Excel. Tiap creator selesai, barisnya langsung
                                ditambahkan ke tab "Hasil WA" di spreadsheet yang sama
                                (tab dibuat otomatis). Tab input TIDAK pernah diubah.
Akses Google lewat service account (credentials.json di folder ini).

Satu-satunya fungsi yang menulis ke Google Sheets: SheetsClient.siapkan_hasil() dan
SheetsClient.tambah_baris(), keduanya hanya menyentuh tab TAB_HASIL.
"""

import json
import re
import time
from itertools import chain, islice
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
CREDENTIALS_FILE = BASE_DIR / 'credentials.json'
SETTINGS_FILE = BASE_DIR / 'pengaturan.json'
TAB_HASIL = 'Hasil WA'

HEADER_NAMES = ('creator name', 'creator', 'nama creator', 'username', 'handle', 'creator id')
SKIP_VALUES = {'', '-', '--', 'summary'}
HEADER_SCAN = 30          # header dicari di 30 baris pertama
COLUMNS = ['Input', 'Handle', 'Nickname', 'Creator ID', 'WhatsApp', 'Link WA', 'Email',
           'Zalo', 'Line', 'Facebook', 'Instagram', 'MCN', 'Followers', 'Status']


# ------------------------------------------------------------- pengaturan ----
def muat_pengaturan():
    try:
        return json.loads(SETTINGS_FILE.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}


def simpan_pengaturan(d):
    try:
        SETTINGS_FILE.write_text(json.dumps(d, indent=2, ensure_ascii=False), encoding='utf-8')
    except OSError:
        pass


def email_service_account():
    try:
        return json.loads(CREDENTIALS_FILE.read_text(encoding='utf-8')).get('client_email')
    except (OSError, ValueError):
        return None


# ------------------------------------------------------------ baca daftar ----
def cari_header(rows):
    """(indeks_baris, indeks_kolom) sel header pertama yang dikenali, atau None."""
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if str(v or '').strip().lower() in HEADER_NAMES:
                return i, j
    return None


def nama_unik(values):
    seen = {}
    for v in values:
        v = str(v).strip() if v is not None else ''
        if v.lower() not in SKIP_VALUES:
            seen.setdefault(v, None)
    return list(seen)


def baca_excel(path):
    """Excel -> (nama unik berurutan, label kolom). Tanpa header dikenali: kolom A."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        head = list(islice(rows, HEADER_SCAN))
        pos = cari_header(head)
        i, j = pos or (-1, 0)
        label = str(head[i][j]).strip() if pos else 'kolom A'
        semua = chain(head, rows)
        return nama_unik((r[j] if j < len(r) else None)
                         for k, r in enumerate(semua) if k > i), label
    finally:
        wb.close()


def parse_link_sheet(link):
    """Link/ID spreadsheet -> (spreadsheet_id, gid atau None)."""
    m = re.search(r'/d/([\w-]+)', link)
    g = re.search(r'gid=(\d+)', link)
    return (m.group(1) if m else link.strip()), (int(g.group(1)) if g else None)


def huruf_kolom(j):
    """0 -> A, 25 -> Z, 26 -> AA."""
    s = ''
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        s = chr(65 + r) + s
    return s


# ------------------------------------------------------------ Google Sheets --
class SheetsClient:
    def __init__(self, credentials=CREDENTIALS_FILE):
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        if not Path(credentials).exists():
            raise FileNotFoundError('credentials.json tidak ketemu. Taruh di folder: %s' % BASE_DIR)
        cred = service_account.Credentials.from_service_account_file(
            str(credentials), scopes=['https://www.googleapis.com/auth/spreadsheets'])
        self.email = cred.service_account_email
        self.api = build('sheets', 'v4', credentials=cred, cache_discovery=False).spreadsheets()

    def _jalankan(self, req):
        from googleapiclient.errors import HttpError
        try:
            return req.execute(num_retries=3)
        except HttpError as e:
            if e.resp.status in (403, 404):
                raise PermissionError(
                    'Spreadsheet tidak bisa diakses (HTTP %d). Share ke %s sebagai Editor, '
                    'lalu coba lagi.' % (e.resp.status, self.email)) from e
            if e.resp.status == 400 and 'not supported' in str(e):
                raise ValueError('File ini .xlsx yang di-upload ke Drive, bukan Google Sheets asli. '
                                 'Buka di Drive lalu File > Save as Google Sheets, atau pakai input Excel.') from e
            raise

    def tabs(self, sid):
        d = self._jalankan(self.api.get(spreadsheetId=sid, fields='sheets.properties(sheetId,title)'))
        return [s['properties'] for s in d['sheets']]

    def baca_input(self, link, tab=''):
        """Balikin (nama unik, label kolom, nama tab input, spreadsheet_id). Hanya BACA."""
        sid, gid = parse_link_sheet(link)
        props = self.tabs(sid)
        if tab.strip():
            pilih = next((p for p in props if p['title'] == tab.strip()), None)
            if pilih is None:
                raise ValueError('Tab "%s" tidak ada. Tab yang ada: %s'
                                 % (tab, ', '.join(p['title'] for p in props)))
        else:
            pilih = next((p for p in props if p['sheetId'] == gid), props[0])
        judul = pilih['title']
        if judul == TAB_HASIL:
            raise ValueError('Tab input tidak boleh bernama "%s" (itu tab hasil).' % TAB_HASIL)
        q = "'%s'" % judul.replace("'", "''")
        head = self._jalankan(self.api.values().get(
            spreadsheetId=sid, range='%s!1:%d' % (q, HEADER_SCAN))).get('values', [])
        i, j = cari_header(head) or (-1, 0)
        label = str(head[i][j]).strip() if i >= 0 else 'kolom A'
        h = huruf_kolom(j)
        kol = self._jalankan(self.api.values().get(
            spreadsheetId=sid, range='%s!%s:%s' % (q, h, h))).get('values', [])
        return nama_unik(r[0] if r else '' for r in kol[i + 1:]), label, judul, sid

    def sudah_ada(self, sid):
        """Nama (kolom Input) yang sudah ada di tab hasil -> dilewati kalau dijalankan ulang."""
        if not any(p['title'] == TAB_HASIL for p in self.tabs(sid)):
            return set()
        v = self._jalankan(self.api.values().get(
            spreadsheetId=sid, range="'%s'!A:A" % TAB_HASIL)).get('values', [])
        return {r[0].strip() for r in v[1:] if r}

    def siapkan_hasil(self, sid):
        """[TULIS] Buat tab TAB_HASIL kalau belum ada, dan isi header kalau masih kosong."""
        if not any(p['title'] == TAB_HASIL for p in self.tabs(sid)):
            self._jalankan(self.api.batchUpdate(spreadsheetId=sid, body={'requests': [
                {'addSheet': {'properties': {'title': TAB_HASIL, 'gridProperties': {
                    'rowCount': 1000, 'columnCount': len(COLUMNS), 'frozenRowCount': 1}}}}]}))
        a1 = self._jalankan(self.api.values().get(
            spreadsheetId=sid, range="'%s'!A1" % TAB_HASIL)).get('values')
        if not a1:
            self._jalankan(self.api.values().update(
                spreadsheetId=sid, range="'%s'!A1" % TAB_HASIL, valueInputOption='RAW',
                body={'values': [COLUMNS]}))

    def tambah_baris(self, sid, row):
        """[TULIS] Tambah satu baris di bawah data terakhir tab TAB_HASIL. RAW: nomor tetap teks."""
        self._jalankan(self.api.values().append(
            spreadsheetId=sid, range="'%s'!A1" % TAB_HASIL, valueInputOption='RAW',
            insertDataOption='INSERT_ROWS',
            body={'values': [[str(row.get(k, '') or '') for k in COLUMNS]]}))


# ---------------------------------------------------------------- olah data --
def normalisasi_wa(v):
    """'0821-6733-7702' -> '6282167337702' (format wa.me). Kosong kalau nggak ada nomor."""
    d = re.sub(r'\D', '', str(v or ''))
    if not d:
        return ''
    if d.startswith('0'):
        return '62' + d[1:]
    if d.startswith('8'):
        return '62' + d
    return d


def olah(kal, nama, cache):
    """Satu nama -> satu baris hasil (dict berkunci COLUMNS)."""
    row, c, saran = {'Input': nama}, {}, None
    if re.fullmatch(r'\d{15,}', nama):
        uid = nama
    else:
        c, saran = kal.cari(nama)
        uid = c.get('creator_uid') if c else None
    if not uid:
        row['Status'] = 'Tidak ketemu' + (' (mirip: %s)' % saran if saran else '')
        return row
    if uid not in cache:
        cache[uid] = kal.kontak(uid)
    d = cache[uid]
    kontak = d.get('creatorContent') or d
    wa = str(kontak.get('whatsapp') or '').strip()
    n = normalisasi_wa(wa)
    row.update({
        'Handle': d.get('handle') or (c or {}).get('creator_handle', ''),
        'Nickname': d.get('nickname') or (c or {}).get('creator_nickname', ''),
        'Creator ID': uid, 'WhatsApp': wa, 'Link WA': 'https://wa.me/' + n if n else '',
        'Email': kontak.get('email', ''), 'Zalo': kontak.get('zalo', ''),
        'Line': kontak.get('line', ''), 'Facebook': kontak.get('facebook', ''),
        'Instagram': kontak.get('ins_id', ''), 'MCN': d.get('mcn_name', ''),
        'Followers': d.get('follower_count', ''),
        'Status': 'OK' if wa else 'Tanpa kontak WA'})
    return row


def simpan_excel(rows, path):
    """Simpan xlsx. Kalau file lagi kebuka di Excel, simpan ke nama bercap waktu."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Creator WA'
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
    for r in rows:
        ws.append([r.get(k, '') for k in COLUMNS])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column_letter in ('D', 'E'):          # ID & WA: teks, biar 0 depan nggak hilang
                c.number_format = '@'
    for i, w in enumerate([22, 22, 24, 22, 16, 30, 28, 14, 14, 20, 18, 16, 11, 26]):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    try:
        wb.save(path)
        return str(path)
    except PermissionError:
        alt = '%s_%s.xlsx' % (Path(path).with_suffix(''), time.strftime('%H%M%S'))
        wb.save(alt)
        return alt


class ExcelSink:
    """Kumpulkan baris, simpan ke Excel tiap SIMPAN_TIAP baris dan di akhir."""
    SIMPAN_TIAP = 25

    def __init__(self, path):
        self.path, self.rows = str(path), []

    def __call__(self, row):
        self.rows.append(row)
        if len(self.rows) % self.SIMPAN_TIAP == 0:
            self.simpan()

    def simpan(self):
        self.path = simpan_excel(self.rows, self.path)   # path berubah kalau jatuh ke nama cadangan
        return self.path


class SheetSink:
    """Tiap baris langsung ditambahkan ke tab TAB_HASIL di spreadsheet input."""

    def __init__(self, client, sid):
        self.client, self.sid = client, sid

    def __call__(self, row):
        self.client.tambah_baris(self.sid, row)


def jalankan(names, kal, sink, log=print, progress=None, stop=None):
    """Proses semua nama. True kalau selesai penuh, False kalau dihentikan."""
    cache = {}
    for i, nama in enumerate(names, 1):
        if stop is not None and stop.is_set():
            log('Dihentikan oleh pengguna.\n')
            return False
        row = olah(kal, nama, cache)
        sink(row)
        log('[%d/%d] %s -> %s %s\n' % (i, len(names), nama, row['Status'], row.get('WhatsApp', '')))
        if progress:
            progress(i, len(names))
    return True


if __name__ == '__main__':
    # ---- self-check tanpa jaringan: python inti.py ----
    import os
    import tempfile
    import threading

    assert normalisasi_wa('082167337702') == '6282167337702'
    assert normalisasi_wa('+62 821-6733-7702') == '6282167337702'
    assert normalisasi_wa('82167337702') == '6282167337702'
    assert normalisasi_wa('') == '' and normalisasi_wa(None) == ''
    assert (huruf_kolom(0), huruf_kolom(25), huruf_kolom(26), huruf_kolom(701)) == ('A', 'Z', 'AA', 'ZZ')
    assert parse_link_sheet('https://docs.google.com/spreadsheets/d/1AbC_-x/edit?gid=123#gid=123') == ('1AbC_-x', 123)
    assert parse_link_sheet(' 1AbC_-x ') == ('1AbC_-x', None)
    # header di baris ke-3 (bukan baris 1), kolom ke-2
    assert cari_header([['Laporan'], [], ['Tgl', 'Creator name']]) == (2, 1)
    assert cari_header([['a'], ['b']]) is None
    assert nama_unik(['a', ' a ', 'Summary', '-', None, 'b']) == ['a', 'b']

    class Kal:      # Kalodata tiruan
        def cari(self, n):
            return ({'creator_uid': '1' * 16, 'creator_handle': n}, None) if n != 'x' else (None, 'xx')
        def kontak(self, uid):
            return {'handle': 'h', 'nickname': 'N', 'creatorContent': {'whatsapp': '0812-111', 'email': ''}}

    got = []
    assert jalankan(['a', 'x', '2' * 16], Kal(), got.append, log=lambda s: None)
    assert [r['Status'] for r in got] == ['OK', 'Tidak ketemu (mirip: xx)', 'OK']
    assert got[0]['Link WA'] == 'https://wa.me/62812111'
    ev = threading.Event(); ev.set()
    assert jalankan(['a'], Kal(), got.append, log=lambda s: None, stop=ev) is False and len(got) == 3

    d = tempfile.mkdtemp()
    f = os.path.join(d, 'in.xlsx')
    wb = Workbook(); ws = wb.active
    ws.append(['Laporan']); ws.append(['Tgl', 'Creator name'])
    for n in ('p', 'q', 'p', 'Summary'):
        ws.append(['d', n])
    wb.save(f)
    assert baca_excel(f) == (['p', 'q'], 'Creator name')
    sink = ExcelSink(os.path.join(d, 'out.xlsx'))
    for r in got:
        sink(r)
    out = sink.simpan()
    assert [c.value for c in load_workbook(out).active[2]][4] == '0812-111'
    print('OK: self-check lolos')
